from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.timetable_views import TimetableViewEntry


_HEADERS = [
    "Session ID",
    "Day",
    "Start Period",
    "End Period",
    "Duration",
    "Subject",
    "Course Code",
    "Course Name",
    "Academic Group Code",
    "Academic Group Name",
    "Faculty",
    "Room",
]


_DAY_ORDER = {
    "MONDAY": 0,
    "TUESDAY": 1,
    "WEDNESDAY": 2,
    "THURSDAY": 3,
    "FRIDAY": 4,
    "SATURDAY": 5,
}


def _entry_sort_key(entry: TimetableViewEntry) -> tuple:
    return (
        _DAY_ORDER.get(entry.day.value, 99),
        entry.start_period,
        entry.course_code,
        entry.academic_group_code,
        entry.session_id,
    )


def _safe_sheet_name(name: str, fallback: str) -> str:
    """Return a valid and reasonably unique Excel worksheet name."""

    invalid_characters = {
        "/": "-",
        "\\": "-",
        "*": "-",
        "?": "-",
        ":": "-",
        "[": "(",
        "]": ")",
    }

    cleaned = name

    for character, replacement in invalid_characters.items():
        cleaned = cleaned.replace(character, replacement)

    cleaned = cleaned.strip() or fallback

    # Excel worksheet names are limited to 31 characters.
    return cleaned[:31]


def _write_sheet(
    worksheet,
    entries: tuple[TimetableViewEntry, ...],
) -> None:
    """Write timetable entries to one worksheet."""

    worksheet.append(_HEADERS)

    for cell in worksheet[1]:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for entry in entries:
        duration = (
            entry.end_period
            - entry.start_period
            + 1
        )

        worksheet.append(
            [
                entry.session_id,
                entry.day.value,
                entry.start_period,
                entry.end_period,
                duration,
                entry.subject_name,
                entry.course_code,
                entry.course_name,
                entry.academic_group_code,
                entry.academic_group_name,
                entry.faculty_name,
                entry.room_name or "ONLINE",
            ]
        )

    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 2:
        worksheet.auto_filter.ref = worksheet.dimensions

    widths = {
        "A": 28,
        "B": 14,
        "C": 14,
        "D": 14,
        "E": 12,
        "F": 32,
        "G": 18,
        "H": 34,
        "I": 28,
        "J": 45,
        "K": 25,
        "L": 15,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:1"


def _create_workbook() -> Workbook:
    """Create a workbook with the default worksheet removed."""

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    return workbook


def export_master_workbook(
    entries: tuple[TimetableViewEntry, ...],
    output_path: str | Path,
) -> Path:
    """Export all timetable entries to one master workbook."""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = _create_workbook()

    worksheet = workbook.create_sheet("Master")

    sorted_entries = tuple(
        sorted(entries, key=_entry_sort_key)
    )

    _write_sheet(worksheet, sorted_entries)

    workbook.save(output_path)

    return output_path


def export_day_workbook(
    entries: tuple[TimetableViewEntry, ...],
    output_path: str | Path,
) -> Path:
    """Export one worksheet per day."""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = _create_workbook()

    entries_by_day: dict[
        str,
        list[TimetableViewEntry],
    ] = defaultdict(list)

    for entry in entries:
        entries_by_day[entry.day.value].append(entry)

    for day_name in _DAY_ORDER:
        day_entries = entries_by_day.get(day_name, [])

        worksheet = workbook.create_sheet(
            _safe_sheet_name(
                day_name.title(),
                day_name.title(),
            )
        )

        sorted_entries = tuple(
            sorted(
                day_entries,
                key=lambda entry: (
                    entry.start_period,
                    entry.course_code,
                    entry.academic_group_code,
                    entry.session_id,
                ),
            )
        )

        _write_sheet(worksheet, sorted_entries)

    workbook.save(output_path)

    return output_path


def export_faculty_workbook(
    entries: tuple[TimetableViewEntry, ...],
    output_path: str | Path,
) -> Path:
    """Export one worksheet per faculty member."""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = _create_workbook()

    entries_by_faculty: dict[
        str,
        list[TimetableViewEntry],
    ] = defaultdict(list)

    for entry in entries:
        entries_by_faculty[entry.faculty_name].append(entry)

    for faculty_name in sorted(entries_by_faculty):
        faculty_entries = entries_by_faculty[faculty_name]

        worksheet = workbook.create_sheet(
            _safe_sheet_name(
                faculty_name,
                "Faculty",
            )
        )

        sorted_entries = tuple(
            sorted(
                faculty_entries,
                key=_entry_sort_key,
            )
        )

        _write_sheet(worksheet, sorted_entries)

    workbook.save(output_path)

    return output_path
